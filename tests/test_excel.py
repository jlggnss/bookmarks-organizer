"""Tests for Excel export and import functionality with native Table and KPIs."""

from datetime import datetime
from pathlib import Path
import pytest
import openpyxl

from src.excel import (
    _format_timestamp,
    _parse_to_timestamp,
    extract_bookmarks_with_paths,
    export_bookmarks_to_excel,
    import_bookmarks_from_excel,
)
from src.models import Bookmark, Folder
from src.parser import extract_all_bookmarks
from src.writer import write_bookmarks


class TestTimestampHelpers:
    def test_format_unix_timestamp(self):
        # 1693701459 is 2023-09-03 00:37:39 UTC (exact time depends on local timezone)
        res = _format_timestamp("1693701459")
        assert "2023-09" in res
        assert ":" in res

    def test_format_empty_or_invalid(self):
        assert _format_timestamp("") == ""
        assert _format_timestamp("not-a-number") == "not-a-number"

    def test_parse_to_timestamp(self):
        dt_str = "2023-09-02 12:00:00"
        ts = _parse_to_timestamp(dt_str)
        assert ts.isdigit()
        assert int(ts) > 1_600_000_000

    def test_parse_already_timestamp(self):
        assert _parse_to_timestamp("1693701459") == "1693701459"
        assert _parse_to_timestamp("") == ""


class TestExcelExportAndImport:
    @pytest.fixture
    def sample_tree(self):
        toolbar = Folder(
            title="Bookmarks bar",
            children=[
                Bookmark(title="Quick Tool", url="https://tool.com", add_date="1690000000"),
                Folder(
                    title="Dev",
                    children=[
                        Bookmark(title="GitHub", url="https://github.com", add_date="1680000000"),
                        Bookmark(title="GitLab", url="https://gitlab.com", add_date="1670000000"),
                    ],
                ),
                Folder(
                    title="News",
                    children=[
                        Bookmark(title="BBC", url="https://bbc.com", add_date="1660000000"),
                        # Duplicate URL to test duplicate detection
                        Bookmark(title="GitHub Duplicate", url="https://github.com", add_date="1650000000"),
                    ],
                ),
            ],
        )
        return Folder(title="Bookmarks", children=[toolbar])

    def test_extract_bookmarks_with_paths(self, sample_tree):
        records = extract_bookmarks_with_paths(sample_tree)
        assert len(records) == 5

        # Check paths
        quick_tool = next(r for r in records if r["title"] == "Quick Tool")
        assert quick_tool["folder_path"] == "Bookmarks bar"

        github = next(r for r in records if r["title"] == "GitHub")
        assert github["folder_path"] == "Bookmarks bar/Dev"

    def test_export_to_excel(self, sample_tree, tmp_path):
        excel_path = tmp_path / "test_bookmarks.xlsx"
        export_bookmarks_to_excel(sample_tree, excel_path)

        assert excel_path.exists()

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        assert ws.title == "Bookmarks"

        # Check Table
        assert "Bookmarks" in ws.tables
        table = ws.tables["Bookmarks"]
        assert table.tableStyleInfo.name == "TableStyleMedium21"

        # Check Row 1 KPIs
        assert ws["B1"].value == "=COUNTA(Bookmarks[Bookmark Title])"
        assert "Oldest:" in ws["F1"].value
        assert ws.freeze_panes == "A3"

        # Check Row 2 Headers
        headers = [cell.value for cell in ws[2]]
        assert "[#]" in headers
        assert "Folder Path" in headers
        assert "Bookmark Title" in headers
        assert "URL" in headers
        assert "Date Added" in headers
        assert "Domain" in headers
        assert "Location" in headers
        assert "Duplicate" in headers
        assert "Status" in headers
        assert "Length of Bookmark Name" in headers
        assert "# Bookmarks On Domain" in headers

        # Check data rows (starts at row 3)
        rows = list(ws.iter_rows(min_row=3, values_only=True))
        assert len(rows) == 5

        # Row 1 index should be 1
        assert rows[0][0] == 1

        # Check toolbar root detection
        quick_tool_row = next(r for r in rows if r[2] == "Quick Tool")
        assert quick_tool_row[1] == "Bookmarks bar"
        assert quick_tool_row[6] == "Toolbar Root"

        # Check folder detection
        dev_row = next(r for r in rows if r[2] == "GitHub")
        assert dev_row[1] == "Bookmarks bar/Dev"
        assert dev_row[6] == "Folder"

        # Check duplicate detection (GitHub appears twice)
        assert "Yes" in dev_row[7]

    def test_import_from_excel_reconstructs_tree(self, sample_tree, tmp_path):
        excel_path = tmp_path / "test_export.xlsx"
        export_bookmarks_to_excel(sample_tree, excel_path)

        # Import back
        reconstructed = import_bookmarks_from_excel(excel_path)
        assert reconstructed.title == "Bookmarks"

        # Check top-level folder
        assert len(reconstructed.children) == 1
        toolbar = reconstructed.children[0]
        assert toolbar.title == "Bookmarks bar"

        # Direct toolbar bookmark
        direct_titles = [c.title for c in toolbar.children if isinstance(c, Bookmark)]
        assert "Quick Tool" in direct_titles

        # Subfolders Dev and News
        sub_folders = {c.title: c for c in toolbar.children if isinstance(c, Folder)}
        assert "Dev" in sub_folders
        assert "News" in sub_folders

        dev_bms = [b.title for b in sub_folders["Dev"].children]
        assert "GitHub" in dev_bms
        assert "GitLab" in dev_bms

    def test_import_skips_marked_deletions(self, sample_tree, tmp_path):
        excel_path = tmp_path / "test_delete.xlsx"
        export_bookmarks_to_excel(sample_tree, excel_path)

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # Find Status and Title columns from Header Row (Row 2)
        headers = [cell.value for cell in ws[2]]
        status_col = headers.index("Status") + 1
        title_col = headers.index("Bookmark Title") + 1

        # Mark "Quick Tool" as Delete and "GitLab" as Drop
        for row in range(3, ws.max_row + 1):
            title = ws.cell(row=row, column=title_col).value
            if title == "Quick Tool":
                ws.cell(row=row, column=status_col, value="Delete")
            elif title == "GitLab":
                ws.cell(row=row, column=status_col, value="drop")

        wb.save(excel_path)

        # Re-import
        reconstructed = import_bookmarks_from_excel(excel_path)
        all_bookmarks, _ = extract_all_bookmarks(reconstructed)
        titles = [b.title for b in all_bookmarks]

        assert "Quick Tool" not in titles
        assert "GitLab" not in titles
        assert "GitHub" in titles
        assert "BBC" in titles


class TestToolbarPreservation:
    def test_preserve_toolbar_option(self):
        toolbar = Folder(
            title="Bookmarks bar",
            children=[
                Bookmark(title="Short Tool", url="https://tool.com"),
                Folder(
                    title="Subfolder",
                    children=[Bookmark(title="Sub Bookmark", url="https://sub.com")],
                ),
            ],
        )
        root = Folder(title="Bookmarks", children=[toolbar])

        # Without preserve_toolbar: both are extracted
        bms_normal, prot_normal = extract_all_bookmarks(root, preserve_toolbar=False)
        assert len(bms_normal) == 2
        assert len(prot_normal) == 0

        # With preserve_toolbar: "Short Tool" stays in protected toolbar folder
        bms_preserved, prot_preserved = extract_all_bookmarks(root, preserve_toolbar=True)
        assert len(bms_preserved) == 1
        assert bms_preserved[0].title == "Sub Bookmark"
        assert len(prot_preserved) == 1
        assert prot_preserved[0].title == "Bookmarks bar"
        assert prot_preserved[0].children[0].title == "Short Tool"
